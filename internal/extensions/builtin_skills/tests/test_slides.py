import os
import tempfile
import unittest

from _support import need_modules, need_programs, run_script, script_path


def make_deck(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly review"
    slide.shapes.add_shape(1, Inches(12.5), Inches(1), Inches(2), Inches(1))
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "Revenue rose"
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    prs.save(path)


@need_modules("pptx")
class Outline(unittest.TestCase):
    def test_dump_lists_shapes_text_notes_and_marks_off_slide_shapes(self):
        with tempfile.TemporaryDirectory() as folder:
            deck = os.path.join(folder, "deck.pptx")
            make_deck(deck)
            result = run_script("slides", "pptx_outline", deck)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn("Quarterly review", result.stdout)
            self.assertIn("Revenue rose", result.stdout)
            self.assertIn("OUTSIDE SLIDE", result.stdout)
            self.assertIn("notes: Speaker notes", result.stdout)

    def test_unreadable_file_exits_two(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "bad.pptx")
            with open(path, "w") as handle:
                handle.write("not a deck")
            self.assertEqual(run_script("slides", "pptx_outline", path).returncode, 2)


@need_modules("pptx", "pypdf", "PIL")
@need_programs("soffice", "pdftoppm")
class Render(unittest.TestCase):
    def test_every_slide_becomes_an_image(self):
        with tempfile.TemporaryDirectory() as folder:
            deck = os.path.join(folder, "deck.pptx")
            make_deck(deck)
            seen = run_script("slides", "render_preview", deck, "--outdir", os.path.join(folder, "p"))
            self.assertEqual(seen.returncode, 0, seen.stdout + seen.stderr)
            self.assertIn("pages: 1", seen.stdout)
            self.assertTrue(os.path.exists(os.path.join(folder, "p", "page-01.png")))


class SharedScript(unittest.TestCase):
    def test_both_skills_ship_the_same_render_script(self):
        with open(script_path("slides", "render_preview.py"), "rb") as a, open(script_path("document-report", "render_preview.py"), "rb") as b:
            self.assertEqual(a.read(), b.read())


if __name__ == "__main__":
    unittest.main()

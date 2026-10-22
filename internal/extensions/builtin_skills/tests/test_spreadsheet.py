import json
import os
import tempfile
import unittest

from _support import need_modules, run_script


@need_modules("openpyxl")
class XlsxSummary(unittest.TestCase):
    def workbook(self, folder, bad=False):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Item", "Qty", "Price"])
        ws.append(["A", 2, 10.5])
        ws.append(["B", 3, 1000.5])
        ws["D1"] = "Total"
        ws["D2"] = "=B2*C2"
        if bad:
            ws["E2"] = "#DIV/0!"
        path = os.path.join(folder, "book.xlsx")
        wb.save(path)
        return path

    def test_structure_dump_and_uncomputed_formulas(self):
        with tempfile.TemporaryDirectory() as folder:
            result = run_script("spreadsheet", "xlsx_summary", self.workbook(folder))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn("header: Item | Qty | Price | Total", result.stdout)
            self.assertIn("Price (C): count=2 sum=1011.0", result.stdout)
            self.assertIn("UNCOMPUTED formulas: 1", result.stdout)

    def test_error_cells_fail(self):
        with tempfile.TemporaryDirectory() as folder:
            result = run_script("spreadsheet", "xlsx_summary", self.workbook(folder, bad=True))
            self.assertEqual(result.returncode, 1)
            self.assertIn("E2=#DIV/0!", result.stdout)

    def test_json_output(self):
        with tempfile.TemporaryDirectory() as folder:
            result = run_script("spreadsheet", "xlsx_summary", self.workbook(folder), "--json")
            self.assertEqual(json.loads(result.stdout)["sheets"][0]["rows"], 3)


if __name__ == "__main__":
    unittest.main()

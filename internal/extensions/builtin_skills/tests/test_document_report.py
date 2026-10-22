import os
import tempfile
import unittest

from _support import need_modules, need_programs, run_script


@need_modules("docx", "pypdf", "PIL", "reportlab")
@need_programs("pdftoppm")
class RenderPreview(unittest.TestCase):
    def make_pdf(self, path):
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        pdf = canvas.Canvas(path, pagesize=letter)
        pdf.drawString(72, 700, "Revenue rose in every quarter.")
        pdf.showPage()
        pdf.drawString(72, 700, "Second page")
        pdf.save()

    def test_pdf_pages_are_rendered_and_counted(self):
        with tempfile.TemporaryDirectory() as folder:
            pdf = os.path.join(folder, "report.pdf")
            self.make_pdf(pdf)
            out = os.path.join(folder, "preview")
            seen = run_script("document-report", "render_preview", pdf, "--outdir", out)
            self.assertEqual(seen.returncode, 0, seen.stdout + seen.stderr)
            self.assertIn("pages: 2", seen.stdout)
            self.assertIn("images in pdf: 0", seen.stdout)
            self.assertTrue(os.path.exists(os.path.join(out, "page-02.png")))
            self.assertTrue(os.path.exists(os.path.join(out, "contact.png")))

    @need_programs("soffice")
    def test_docx_is_converted_and_rendered(self):
        import docx
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "report.docx")
            document = docx.Document()
            document.add_heading("Summary", 1)
            document.add_paragraph("Revenue rose.")
            document.save(path)
            seen = run_script("document-report", "render_preview", path, "--outdir", os.path.join(folder, "p"))
            self.assertEqual(seen.returncode, 0, seen.stdout + seen.stderr)
            self.assertIn("pages: 1", seen.stdout)


if __name__ == "__main__":
    unittest.main()

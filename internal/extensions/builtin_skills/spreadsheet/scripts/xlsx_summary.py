#!/usr/bin/env python3
"""Summarise an XLSX workbook so its content can be verified.

Reports, per sheet: used range, non-empty row count, header row, formula count,
error cells, and for every column holding numbers the count and sum.

Formulas written by openpyxl carry no computed value until a spreadsheet
program recalculates them. --recalc runs LibreOffice headless to recalculate a
copy first, so sums include formula results; without it, formula cells that
have no cached value are counted and listed as uncomputed.

Exit status 1 when error cells exist.
Usage: xlsx_summary.py book.xlsx [--recalc] [--json]
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "Err:")


def recalculated_copy(path, workdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("soffice is not available for --recalc")
    source = os.path.join(workdir, "in", os.path.basename(path))
    os.makedirs(os.path.dirname(source))
    shutil.copy(path, source)
    out = os.path.join(workdir, "out")
    os.makedirs(out)
    profile = "file://" + os.path.join(workdir, "profile")
    subprocess.run([soffice, "-env:UserInstallation=" + profile, "--headless", "--calc",
                    "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, source],
                   check=True, capture_output=True, timeout=180)
    result = os.path.join(out, os.path.basename(path))
    if not os.path.exists(result):
        raise RuntimeError("recalculation produced no file")
    return result


def summarise(path, recalc=False):
    import openpyxl
    workdir = tempfile.mkdtemp(prefix="xlsx-summary-")
    try:
        values_path = recalculated_copy(path, workdir) if recalc else path
        formulas = openpyxl.load_workbook(path, data_only=False)
        values = openpyxl.load_workbook(values_path, data_only=True)
        sheets = []
        for ws in formulas.worksheets:
            vs = values[ws.title] if ws.title in values.sheetnames else None
            info = {"name": ws.title, "range": ws.dimensions, "rows": 0, "formulas": 0,
                    "uncomputed_formulas": 0, "error_cells": [], "header": [], "columns": {}}
            header_row = None
            for row in ws.iter_rows():
                if any(c.value not in (None, "") for c in row):
                    info["rows"] += 1
                    if header_row is None:
                        header_row = row
                        info["header"] = [str(c.value) if c.value is not None else "" for c in row]
            for row in ws.iter_rows():
                for cell in row:
                    raw = cell.value
                    is_formula = isinstance(raw, str) and raw.startswith("=")
                    shown = vs[cell.coordinate].value if vs is not None else raw
                    if is_formula:
                        info["formulas"] += 1
                        if shown is None:
                            info["uncomputed_formulas"] += 1
                    if isinstance(shown, str) and shown.startswith(ERROR_TOKENS):
                        info["error_cells"].append("%s=%s" % (cell.coordinate, shown))
                    if isinstance(shown, bool) or not isinstance(shown, (int, float)):
                        continue
                    if header_row is not None and cell.row == header_row[0].row:
                        continue
                    label = info["header"][cell.column - 1] if cell.column - 1 < len(info["header"]) else ""
                    key = "%s (%s)" % (label or "column", cell.column_letter)
                    entry = info["columns"].setdefault(key, {"count": 0, "sum": Decimal(0)})
                    entry["count"] += 1
                    entry["sum"] += Decimal(str(shown))
            for entry in info["columns"].values():
                entry["sum"] = float(entry["sum"])
            sheets.append(info)
        return {"file": os.path.basename(path), "recalculated": bool(recalc), "sheets": sheets}
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def render(report):
    lines = ["%s (recalculated: %s)" % (report["file"], "yes" if report["recalculated"] else "no"),
             "sheets: %d" % len(report["sheets"])]
    for info in report["sheets"]:
        lines.append("")
        lines.append("sheet %r range %s, %d non-empty rows, %d formulas" % (info["name"], info["range"], info["rows"], info["formulas"]))
        lines.append("  header: %s" % " | ".join(info["header"]))
        if info["uncomputed_formulas"]:
            lines.append("  UNCOMPUTED formulas: %d (run with --recalc)" % info["uncomputed_formulas"])
        if info["error_cells"]:
            lines.append("  ERROR cells: %s" % ", ".join(info["error_cells"][:20]))
        for key, entry in info["columns"].items():
            lines.append("  numeric %s: count=%d sum=%.4f" % (key, entry["count"], entry["sum"]))
    return "\n".join(lines)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("path")
    parser.add_argument("--recalc", action="store_true")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        report = summarise(args.path, args.recalc)
    except ImportError:
        print("openpyxl is not installed in this Python", file=sys.stderr)
        return 2
    except Exception as exc:  # unreadable workbook, missing soffice
        print("cannot summarise: %s" % exc, file=sys.stderr)
        return 2
    print(json.dumps(report, indent=2) if args.json else render(report))
    failures = []
    for info in report["sheets"]:
        if info["error_cells"]:
            failures.append("sheet %s has error cells" % info["name"])
    for text in failures:
        print("FAIL: " + text)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())

#!/usr/bin/env python3
"""Check a PPTX deck: bounds, text fit, titles, fonts, overlaps and counts.

Errors (exit status 1): a shape extending beyond the slide, estimated text
overflow of its box, a slide without a title, a failed --expect-* assertion.
Warnings: text above 90% of its box, fonts under 12 pt (10 pt in tables),
overlapping shapes, a slide without content.

Text height is estimated from font size, box width and glyph widths (see
deckkit.py); it screens for problems but the rendered slide decides.
Prints per-slide titles and counts of charts, pictures and tables.
Usage: check_deck.py deck.pptx [--expect-slides N] [--expect-charts N]
       [--expect-pictures N] [--expect-tables N] [--json]
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import deckkit  # noqa: E402

EMU = 914400.0
TOLERANCE = 0.02  # inches of slack on the bounds


def inches(value):
    return None if value is None else value / EMU


def run_sizes(frame, default):
    sizes = []
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return sizes or [default]


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # group
            for inner in walk(shape.shapes):
                yield inner


def overlap_ratio(a, b):
    left, top = max(a[0], b[0]), max(a[1], b[1])
    right, bottom = min(a[0] + a[2], b[0] + b[2]), min(a[1] + a[3], b[1] + b[3])
    if right <= left or bottom <= top:
        return 0.0
    area = (right - left) * (bottom - top)
    return area / max(min(a[2] * a[3], b[2] * b[3]), 1e-9)


def check(path):
    from pptx import Presentation
    prs = Presentation(path)
    slide_w, slide_h = inches(prs.slide_width), inches(prs.slide_height)
    report = {"file": os.path.basename(path), "slide_size_in": [round(slide_w, 3), round(slide_h, 3)], "slides": [], "errors": [], "warnings": []}
    totals = {"charts": 0, "pictures": 0, "tables": 0}
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        info = {"number": number, "title": title, "charts": 0, "pictures": 0, "tables": 0, "text_shapes": 0}
        if not title:
            report["errors"].append("slide %d: no title" % number)
        boxes = []
        for shape in walk(slide.shapes):
            left, top, width, height = inches(shape.left), inches(shape.top), inches(shape.width), inches(shape.height)
            name = shape.name
            if None in (left, top, width, height):
                continue
            if left < -TOLERANCE or top < -TOLERANCE or left + width > slide_w + TOLERANCE or top + height > slide_h + TOLERANCE:
                report["errors"].append("slide %d: %r extends beyond the slide (left %.2f top %.2f right %.2f bottom %.2f in; slide %.2f x %.2f)" % (
                    number, name, left, top, left + width, top + height, slide_w, slide_h))
            if shape.has_chart if hasattr(shape, "has_chart") else False:
                info["charts"] += 1
            elif shape.shape_type == 13:
                info["pictures"] += 1
            elif getattr(shape, "has_table", False) and shape.has_table:
                info["tables"] += 1
                table = shape.table
                needed = 0.0
                for row in table.rows:
                    row_need = 0.0
                    for c, cell in enumerate(row.cells):
                        sizes = run_sizes(cell.text_frame, 18)
                        if min(sizes) < 10:
                            report["warnings"].append("slide %d: table text at %.0f pt" % (number, min(sizes)))
                        row_need = max(row_need, deckkit.text_height_in([cell.text_frame.text], inches(table.columns[c].width), max(sizes), margin_in=0.1))
                    needed += max(row_need, inches(row.height) or 0)
                if top + needed > slide_h + TOLERANCE:
                    report["errors"].append("slide %d: table %r likely grows to %.1f in tall and runs off the slide" % (number, name, needed))
            if shape.has_text_frame and shape.text_frame.text.strip():
                info["text_shapes"] += 1
                frame = shape.text_frame
                is_title = shape == title_shape
                sizes = run_sizes(frame, 44 if is_title else 18)
                paragraphs = [p.text for p in frame.paragraphs]
                estimate = deckkit.text_height_in(paragraphs, width, max(sizes))
                if estimate > height * 1.05:
                    report["errors"].append("slide %d: text in %r likely overflows its box (needs about %.2f in, box is %.2f in)" % (number, name, estimate, height))
                elif estimate > height * 0.9:
                    report["warnings"].append("slide %d: text in %r fills %.0f%% of its box" % (number, name, 100 * estimate / height))
                if min(sizes) < 12:
                    report["warnings"].append("slide %d: %r uses %.0f pt text" % (number, name, min(sizes)))
            if shape.has_text_frame and shape.text_frame.text.strip() or shape.shape_type == 13 or getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
                boxes.append((left, top, width, height, name))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                if overlap_ratio(boxes[i][:4], boxes[j][:4]) > 0.15:
                    report["warnings"].append("slide %d: %r overlaps %r" % (number, boxes[i][4], boxes[j][4]))
        if not (info["charts"] or info["pictures"] or info["tables"] or info["text_shapes"] > 1):
            report["warnings"].append("slide %d has little content" % number)
        for key in totals:
            totals[key] += info[key]
        report["slides"].append(info)
    report["totals"] = dict(totals, slides=len(report["slides"]))
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("path")
    parser.add_argument("--expect-slides", type=int)
    parser.add_argument("--expect-charts", type=int)
    parser.add_argument("--expect-pictures", type=int)
    parser.add_argument("--expect-tables", type=int)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        report = check(args.path)
    except ImportError:
        print("python-pptx is not installed in this Python", file=sys.stderr)
        return 2
    totals = report["totals"]
    for key, wanted in (("slides", args.expect_slides), ("charts", args.expect_charts), ("pictures", args.expect_pictures), ("tables", args.expect_tables)):
        if wanted is not None and totals[key] != wanted:
            report["errors"].append("expected %d %s, found %d" % (wanted, key, totals[key]))
    if args.json:
        print(json.dumps(report, indent=2))
    else:
        print("%s: %d slides (%.2f x %.2f in), %d charts, %d pictures, %d tables" % (
            report["file"], totals["slides"], report["slide_size_in"][0], report["slide_size_in"][1], totals["charts"], totals["pictures"], totals["tables"]))
        for info in report["slides"]:
            print("  %2d. %s [charts %d, pictures %d, tables %d]" % (info["number"], info["title"] or "(no title)", info["charts"], info["pictures"], info["tables"]))
        for text in report["warnings"]:
            print("WARNING: " + text)
        for text in report["errors"]:
            print("ERROR: " + text)
        if not report["errors"]:
            print("no errors")
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())

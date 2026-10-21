#!/usr/bin/env python3
"""Build a 16:9 PPTX deck from a JSON spec with python-pptx.

Spec (JSON):
  {"slides": [
    {"layout": "title",   "title": "...", "subtitle": "..."},
    {"layout": "bullets", "title": "...", "bullets": ["...", "..."], "notes": "..."},
    {"layout": "chart",   "title": "...", "chart": {"kind": "bar|barh|line|pie",
        "categories": ["Q1", "Q2"], "series": [{"name": "Revenue", "values": [1, 2]}]},
        "bullets": ["optional side points"]},
    {"layout": "table",   "title": "...", "columns": ["A", "B"], "rows": [["x", 1]]},
    {"layout": "image",   "title": "...", "image": "figure.png", "caption": "..."},
    {"layout": "two_column", "title": "...",
        "left": {"heading": "...", "bullets": ["..."]}, "right": {"heading": "...", "bullets": ["..."]}}]}

Every slide needs a title. Text sizes are chosen to fit the estimated box
height (see deckkit.py); a slide whose text cannot fit even at the smallest
size stops the build so the content can be split across slides.
Usage: build_deck.py spec.json out.pptx
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import deckkit  # noqa: E402

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
BODY_TOP = 1.65
BODY_BOTTOM = SLIDE_H - 0.6
ACCENT = (0x30, 0x54, 0x96)
INK = (0x22, 0x22, 0x22)


class SpecError(Exception):
    pass


def numeric(values, where):
    out = []
    for value in values:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise SpecError("%s: %r is not a number" % (where, value))
        out.append(float(value))
    return out


def set_text(frame, paragraphs, size, bold=False, color=INK, align=None, bullets=False):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    frame.word_wrap = True
    for index, text in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = ("• " + text) if bullets else text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)
        if align is not None:
            paragraph.alignment = align
        if bullets:
            paragraph.space_after = Pt(size * 0.35)


def place_title(slide, text, big=False):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches
    shape = slide.shapes.title
    width = SLIDE_W - 2 * MARGIN
    height = 1.1
    if big:
        shape.left, shape.top, shape.width, shape.height = Inches(MARGIN), Inches(2.2), Inches(width), Inches(1.6)
        height = 1.6
    else:
        shape.left, shape.top, shape.width, shape.height = Inches(MARGIN), Inches(0.4), Inches(width), Inches(height)
    size = deckkit.fit_size([text], width, height, 48 if big else 34, 22)
    if deckkit.text_height_in([text], width, size) > height:
        raise SpecError("title too long to fit: %r" % text)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(frame, [text], size, bold=True, color=ACCENT, align=PP_ALIGN.LEFT)


def add_box(slide, left, top, width, height, paragraphs, start, minimum, bullets=False, bold=False, label=""):
    from pptx.util import Inches
    size = deckkit.fit_size(paragraphs, width, height, start, minimum)
    if deckkit.text_height_in(paragraphs, width, size) > height:
        raise SpecError("text does not fit%s even at %d pt; split it across slides" % (" in " + label if label else "", minimum))
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    set_text(box.text_frame, paragraphs, size, bullets=bullets, bold=bold)
    return box


def add_chart(slide, spec, left, top, width, height):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "barh": XL_CHART_TYPE.BAR_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
    kind = spec.get("kind", "bar")
    if kind not in kinds:
        raise SpecError("unknown chart kind %r" % kind)
    categories = [str(c) for c in spec.get("categories") or []]
    series = spec.get("series") or []
    if not categories or not series:
        raise SpecError("chart needs categories and series")
    data = CategoryChartData()
    data.categories = categories
    for item in series:
        values = numeric(item.get("values") or [], "chart series %r" % item.get("name"))
        if len(values) != len(categories):
            raise SpecError("series %r has %d values for %d categories" % (item.get("name"), len(values), len(categories)))
        data.add_series(str(item.get("name", "")), values)
    frame = slide.shapes.add_chart(kinds[kind], Inches(left), Inches(top), Inches(width), Inches(height), data)
    chart = frame.chart
    chart.font.size = Pt(14)
    chart.has_legend = kind == "pie" or len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if kind == "pie":
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.number_format = "0%"
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
    return frame


def add_table(slide, columns, rows, left, top, width, height):
    from pptx.util import Inches, Pt
    if not columns:
        raise SpecError("table needs columns")
    count = len(rows) + 1
    size = 16 if count <= 6 else 14 if count <= 9 else 12 if count <= 12 else 10
    row_h = height / count
    needed = deckkit.text_height_in(["x"], width / len(columns), size, margin_in=0.1) + 0.05
    if needed > row_h:
        raise SpecError("table has too many rows for one slide (%d rows); split it" % len(rows))
    frame = slide.shapes.add_table(count, len(columns), Inches(left), Inches(top), Inches(width), Inches(row_h * count))
    table = frame.table
    for c, name in enumerate(columns):
        table.cell(0, c).text = str(name)
    for r, row in enumerate(rows, start=1):
        for c in range(len(columns)):
            value = row[c] if c < len(row) else ""
            table.cell(r, c).text = "" if value is None else str(value)
    for r in range(count):
        table.rows[r].height = Inches(row_h)
        for c in range(len(columns)):
            for paragraph in table.cell(r, c).text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)
    return frame


def add_picture(slide, path, left, top, width, height):
    from PIL import Image
    from pptx.util import Inches
    if not os.path.exists(path):
        raise SpecError("image file missing: %s" % path)
    with Image.open(path) as probe:
        w, h = probe.size
    scale = min(width / w, height / h)
    pw, ph = w * scale, h * scale
    return slide.shapes.add_picture(path, Inches(left + (width - pw) / 2), Inches(top + (height - ph) / 2), Inches(pw), Inches(ph))


def build(spec, out):
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    layout = prs.slide_layouts[5]  # title only: a real title placeholder on every slide
    width = SLIDE_W - 2 * MARGIN
    body_h = BODY_BOTTOM - BODY_TOP
    slides = spec.get("slides") or []
    if not slides:
        raise SpecError("spec has no slides")
    for number, item in enumerate(slides, start=1):
        title = str(item.get("title") or "").strip()
        if not title:
            raise SpecError("slide %d has no title" % number)
        kind = item.get("layout", "bullets")
        slide = prs.slides.add_slide(layout)
        try:
            if kind == "title":
                place_title(slide, title, big=True)
                if item.get("subtitle"):
                    add_box(slide, MARGIN, 4.0, width, 1.2, [str(item["subtitle"])], 26, 16, label="subtitle")
            else:
                place_title(slide, title)
                if kind == "bullets":
                    add_box(slide, MARGIN, BODY_TOP, width, body_h, [str(b) for b in item.get("bullets") or []], 28, 16, bullets=True, label="the body")
                elif kind == "chart":
                    side = [str(b) for b in item.get("bullets") or []]
                    chart_w = width * (0.58 if side else 1.0)
                    add_chart(slide, item.get("chart") or {}, MARGIN, BODY_TOP, chart_w, body_h)
                    if side:
                        add_box(slide, MARGIN + chart_w + 0.3, BODY_TOP, width - chart_w - 0.3, body_h, side, 22, 14, bullets=True, label="the side points")
                elif kind == "table":
                    add_table(slide, item.get("columns") or [], item.get("rows") or [], MARGIN, BODY_TOP, width, min(body_h, 0.55 * (len(item.get("rows") or []) + 1)))
                elif kind == "image":
                    cap = [str(item["caption"])] if item.get("caption") else []
                    add_picture(slide, str(item.get("image", "")), MARGIN, BODY_TOP, width, body_h - (0.6 if cap else 0))
                    if cap:
                        add_box(slide, MARGIN, BODY_BOTTOM - 0.5, width, 0.5, cap, 16, 12, label="the caption")
                elif kind == "two_column":
                    half = (width - 0.4) / 2
                    for index, side_key in enumerate(("left", "right")):
                        side = item.get(side_key) or {}
                        left = MARGIN + index * (half + 0.4)
                        top = BODY_TOP
                        if side.get("heading"):
                            add_box(slide, left, top, half, 0.6, [str(side["heading"])], 24, 16, bold=True, label="a column heading")
                            top += 0.7
                        add_box(slide, left, top, half, BODY_BOTTOM - top, [str(b) for b in side.get("bullets") or []], 24, 14, bullets=True, label="a column")
                else:
                    raise SpecError("unknown layout %r" % kind)
            if item.get("notes"):
                slide.notes_slide.notes_text_frame.text = str(item["notes"])
        except SpecError as exc:
            raise SpecError("slide %d (%s): %s" % (number, title, exc))
    prs.save(out)
    return len(slides)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spec")
    parser.add_argument("out")
    args = parser.parse_args(argv)
    with open(args.spec, encoding="utf-8") as handle:
        spec = json.load(handle)
    try:
        count = build(spec, args.out)
    except SpecError as exc:
        print("spec error: %s" % exc, file=sys.stderr)
        return 1
    except ImportError as exc:
        print("missing library: %s" % exc, file=sys.stderr)
        return 2
    print("wrote %s with %d slide(s)" % (args.out, count))
    return 0


if __name__ == "__main__":
    sys.exit(main())

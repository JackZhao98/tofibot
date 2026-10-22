#!/usr/bin/env python3
"""Print what a PPTX file holds: slide size and, per slide, every shape.

For each shape: kind (text, picture, chart, table, other), position and size in
inches, whether any part lies outside the slide, its text, and the speaker
notes. A plain dump for reading the file back; it judges nothing.
Usage: pptx_outline.py deck.pptx [--json]
"""
from __future__ import annotations

import argparse
import json
import os
import sys

EMU = 914400.0


def kind(shape):
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if shape.shape_type == 13:
        return "picture"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return "other"


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from walk(shape.shapes)


def outline(path):
    from pptx import Presentation
    prs = Presentation(path)
    width, height = prs.slide_width / EMU, prs.slide_height / EMU
    result = {"file": os.path.basename(path), "slide_size_in": [round(width, 2), round(height, 2)], "slides": []}
    for number, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in walk(slide.shapes):
            if None in (shape.left, shape.top, shape.width, shape.height):
                continue
            left, top, w, h = shape.left / EMU, shape.top / EMU, shape.width / EMU, shape.height / EMU
            info = {"name": shape.name, "kind": kind(shape), "box_in": [round(v, 2) for v in (left, top, w, h)],
                    "outside_slide": left < -0.02 or top < -0.02 or left + w > width + 0.02 or top + h > height + 0.02}
            if info["kind"] == "text":
                info["text"] = shape.text_frame.text
            if info["kind"] == "table":
                info["rows"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
            if info["kind"] == "chart":
                info["chart_type"] = str(shape.chart.chart_type)
            shapes.append(info)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        result["slides"].append({"number": number, "shapes": shapes, "notes": notes})
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("path")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)
    try:
        data = outline(args.path)
    except ImportError:
        print("python-pptx is not installed in this Python", file=sys.stderr)
        return 2
    except Exception as exc:
        print("cannot read: %s" % exc, file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(data, indent=2))
        return 0
    print("%s: %d slides, %.2f x %.2f in" % (data["file"], len(data["slides"]), *data["slide_size_in"]))
    for slide in data["slides"]:
        print("slide %d" % slide["number"])
        for s in slide["shapes"]:
            flag = " OUTSIDE SLIDE" if s["outside_slide"] else ""
            print("  %s %r at %s%s" % (s["kind"], s["name"], s["box_in"], flag))
            for line in (s.get("text") or "").splitlines():
                print("    | " + line)
            for row in s.get("rows", []):
                print("    | " + " | ".join(row))
        if slide["notes"]:
            print("  notes: " + slide["notes"].replace("\n", " "))
    return 0


if __name__ == "__main__":
    sys.exit(main())
